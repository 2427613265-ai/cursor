#!/usr/bin/env python3
"""Formal / numbering / 所述-origin checks for the revised claims."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

DOC = Path("/workspace/cases/PA26-1238-I-A/权利要求书-修订稿.docx")
LOG = Path("/opt/cursor/artifacts/claims_validation.log")


def paras(path: Path) -> list[str]:
    d = Document(str(path))
    return [p.text.strip() for p in d.paragraphs if p.text.strip()]


def split_claims(texts: list[str]) -> dict[int, str]:
    claims: dict[int, list[str]] = {}
    cur = None
    for t in texts:
        m = re.match(r"^(\d+)\.", t)
        if m:
            cur = int(m.group(1))
            claims[cur] = [t]
        elif cur is not None:
            claims[cur].append(t)
    return {k: "\n".join(v) for k, v in claims.items()}


def cited(text: str) -> list[int]:
    head = text.split("其特征在于")[0]
    nums = []
    for m in re.finditer(r"权利要求(\d+)至(\d+)", head):
        nums.extend(range(int(m.group(1)), int(m.group(2)) + 1))
    for m in re.finditer(r"权利要求(\d+)", head):
        n = int(m.group(1))
        if n not in nums:
            nums.append(n)
    return sorted(set(nums))


def main() -> int:
    lines = paras(DOC)
    claims = split_claims(lines)
    errors: list[str] = []
    notes: list[str] = []

    notes.append(f"claims found: {sorted(claims)}")
    if sorted(claims) != list(range(1, 11)):
        errors.append(f"编号不连续: {sorted(claims)}")

    for n, body in claims.items():
        if n == 1:
            continue
        cs = cited(body)
        if any(c >= n for c in cs):
            errors.append(f"权{n}引用了不在前的权利要求: {cs}")
        if "任一项" in body.split("其特征在于")[0] or "或" in body.split("其特征在于")[0]:
            notes.append(f"权{n} 可能是多项从权, cited={cs}")

    c1 = claims[1]
    for token in [
        "L1",
        "R1",
        "J1",
        "对称节段对",
        "隔仓浇筑",
        "回填浇筑",
        "被跳过的对称节段对",
        "循环进行隔仓浇筑",
        "其余非吊索区湿接缝",
        "绝对值不小于2",
        "分批次",
    ]:
        if token not in c1:
            errors.append(f"权1缺少关键用语: {token}")

    if "根据权利要求1至9" in "".join(claims.values()):
        errors.append("仍存在引用1至9")

    if "判断所述第一阶段的湿接缝混凝土是否达到预设强度，若是" in c1:
        errors.append("权1 S4仍用若是双分支")

    # 所述 origin: crude
    first_intro: dict[str, int] = {}
    said_pat = re.compile(r"所述([\u4e00-\u9fffA-Za-z0-9₀-₉′_ΔδδηηWLkRjp]{1,20})")
    for n in sorted(claims):
        # strip 根据权利要求X所述的主题
        body = re.sub(r"^.*?其特征在于，", "", claims[n], count=1, flags=re.S)
        # introductions without 所述 in claim 1 S2 names
        if n == 1:
            for name in [
                "钢主梁节段",
                "吊索",
                "连续钢梁骨架",
                "预制混凝土桥面板",
                "湿接缝",
                "吊索节段",
                "对称节段对",
                "吊索区湿接缝",
                "非吊索区湿接缝",
                "预设跨中区段",
                "浇筑批次",
                "跳仓",
                "预设强度",
            ]:
                first_intro.setdefault(name, n)
        for name in [
            "初始主缆水平张力H₀",
            "当前主缆水平张力H₁",
            "当前竖向位移δ",
            "影响系数",
            "预设指标",
            "控制偏差Δh",
            "竖向位移监测点",
            "主缆变形包络线宽度W_env",
            "预设阈值",
            "变形极值测点",
        ]:
            if name in claims[n] and name not in first_intro:
                # only if not 所述-only first
                pass
        for m in said_pat.finditer(body):
            term = m.group(1)
            # skip theme
            if term.startswith("悬索桥"):
                continue

    # chain coverage for key 所述 terms
    chains = {
        2: [2, 1],
        3: [3, 1],
        4: [4, 1],
        5: [5, 1],
        6: [6, 1],
        7: [7, 6, 1],
        8: [8, 1],
        9: [9, 1],
        10: [10, 1],
    }
    term_first = {
        "S3": 1,
        "预设跨中区段": 1,
        "非吊索区湿接缝": 1,
        "湿接缝": 1,
        "对称节段对": 1,
        "吊索节段": 1,
        "第二阶段": 1,
        "第一阶段": 1,
        "预设范围": 1,
        "主缆": 1,
        "预设阈值": 6,
        "W_env": 6,
        "湿接缝混凝土": 1,
        "预制混凝土桥面板": 1,
    }
    for n, chain in chains.items():
        body = claims[n]
        for term, first in term_first.items():
            if f"所述{term}" in body or f"所述{term}" in body.replace("所述S3", "所述S3"):
                if first not in chain and f"所述{term}" in body:
                    # W_env written as 所述W_env
                    if term == "W_env" and "所述W_env" not in body:
                        continue
                    if f"所述{term}" in body and first not in chain:
                        errors.append(f"权{n} 所述{term} 首次在权{first}，链{chain}未覆盖")

    c3 = claims[3]
    for token in [
        "隔仓浇筑的每一浇筑批次开始前",
        "单位浇筑方量",
        "预先确定的计划浇筑方量",
        "候选对称节段对",
        "作为上一浇筑批次",
        "本轮开始时尚未浇筑",
        "所述本批次浇筑对象",
        "所述被跳过的对称节段对中尚未浇筑",
        "在所述回填浇筑完成后仍存在尚未浇筑的其余非吊索区湿接缝的情况下",
    ]:
        if token not in c3:
            errors.append(f"权3缺少质检改写用语: {token}")
    for bad in [
        "单位浇筑荷载",
        "该对称节段对",
        "分多个浇筑批次",
        "对称浇筑其中尚未浇筑",
        "S36所述编号差约束",
        "S37、",
        "S38、",
        "S39、",
        "如果",
        "当……时",
        "编号之差",
    ]:
        if bad in c3:
            errors.append(f"权3不应再出现: {bad}")
    if "S31、隔仓浇筑的每一浇筑批次" not in c3:
        errors.append("权3子步未在本条内从S31起编")

    # specific
    if "所述预设阈值" in claims[7] and 6 not in chains[7]:
        errors.append("权7预设阈值链错误")
    if "所述预设范围" in claims[9] and 1 not in chains[9]:
        errors.append("权9预设范围链错误")

    notes.append("权1 length chars: %d" % len(c1))
    c1_compact = "".join(c1.split())
    notes.append("权1 compact chars: %d" % len(c1_compact))
    if len(c1_compact) > 500:
        errors.append(f"权1超过500字: {len(c1_compact)}")
    notes.append("abstract will be checked from full docx")

    full = Path("/workspace/cases/PA26-1238-I-A/一种悬索桥钢混组合梁原位叠合浇筑施工方法-权利要求修订稿.docx")
    if full.exists():
        fp = Document(str(full)).paragraphs
        abs_txt = fp[2].text
        notes.append(f"摘要字数(含标点): {len(abs_txt)}")
        if len(abs_txt) > 300:
            errors.append(f"摘要超过300字: {len(abs_txt)}")
        if "L1" not in fp[15].text:
            errors.append("全文稿权1 S2未写入L1编号")
        if "绝对值不小于2" not in fp[16].text:
            errors.append("全文稿权1 S3未写入跳仓编号差")

    log = "\n".join(["NOTES:"] + notes + ["", "ERRORS:"] + (errors or ["none"])) + "\n"
    LOG.write_text(log, encoding="utf-8")
    print(log)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
