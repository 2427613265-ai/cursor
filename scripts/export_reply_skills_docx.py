#!/usr/bin/env python3
"""把「审查意见答复」全部 skill 正文导出为一份 Word 文档。

默认写入：
  .cursor/skills/专利/审查意见答复/创造性答复撰写skill汇编.docx

用法：
  python3 scripts/export_reply_skills_docx.py
  python3 scripts/export_reply_skills_docx.py --output /path/to/out.docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parents[1]
REPLY_ROOT = ROOT / ".cursor" / "skills" / "专利" / "审查意见答复"
DEFAULT_OUTPUT = REPLY_ROOT / "创造性答复撰写skill汇编.docx"

NAVY = "1F4E79"
NAVY_RGB = RGBColor(0x1F, 0x4E, 0x79)
RULE = "D5DEE8"
QUOTE_BG = "F4F7FB"
CODE_BG = "F3F4F6"
LABEL_BG = "E8EEF4"
ZEBRA = "F7F9FB"

# 跑案顺序。新增 SKILL.md 时必须补进本表，脚本会据此核对目录，避免漏导出。
SECTIONS: list[dict[str, str]] = [
    {
        "path": "README.md",
        "title": "审查意见答复",
        "role": "分支说明",
        "kind": "说明",
    },
    {
        "path": "创造性答复/README.md",
        "title": "创造性答复",
        "role": "目录说明、挂接契约与迁出映射",
        "kind": "说明",
    },
    {
        "path": "创造性答复/创造性答复总控/SKILL.md",
        "title": "创造性答复总控",
        "role": "入口：一通/二通、plan、分步挂接、其他权利要求、组装",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复共用硬性规则/SKILL.md",
        "title": "创造性答复共用硬性规则",
        "role": "命题权属、语言分轨、两种协同、加粗与流程禁令",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复反驳手法工具箱/SKILL.md",
        "title": "创造性答复反驳手法工具箱",
        "role": "第1步选法（手法 A–I）",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复修改策略/SKILL.md",
        "title": "创造性答复修改策略",
        "role": "第2步：情形分支、修改分层，与 plan 同批",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复修改说明撰写/SKILL.md",
        "title": "创造性答复修改说明撰写",
        "role": "第3步：修改说明 / 坚持原权",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复区别特征与技术问题撰写/SKILL.md",
        "title": "创造性答复区别特征与技术问题撰写",
        "role": "2.1 / 3.1",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复非显而易见性撰写/SKILL.md",
        "title": "创造性答复非显而易见性撰写",
        "role": "2.2 / 3.2 / 二通第一部分",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复显著技术进步撰写/SKILL.md",
        "title": "创造性答复显著技术进步撰写",
        "role": "2.3 / 3.3",
        "kind": "skill",
    },
    {
        "path": "创造性答复/创造性答复意见陈述书撰写/SKILL.md",
        "title": "创造性答复意见陈述书撰写",
        "role": "旧名跳转页（改走总控）",
        "kind": "skill",
    },
]

CN_NUM = "一二三四五六七八九十"

FRONTMATTER_RE = re.compile(r"\A---\n(.*?)\n---\n?", re.DOTALL)
LIST_RE = re.compile(r"^(\s*)([-*+]|\d+\.)\s+(.*)$")
CHECK_RE = re.compile(r"^(\s*)-\s+\[([ xX])\]\s+(.*)$")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
SEP_CELL_RE = re.compile(r"^:?-{3,}:?$")


def cn_index(n: int) -> str:
    if 1 <= n <= 10:
        return CN_NUM[n - 1]
    if n < 20:
        return "十" + CN_NUM[n - 11]
    raise ValueError(n)


def parse_frontmatter(text: str) -> tuple[dict[str, str], str]:
    match = FRONTMATTER_RE.match(text)
    if not match:
        return {}, text
    meta: dict[str, str] = {}
    for raw in match.group(1).splitlines():
        if not raw.strip() or raw.strip().startswith("#"):
            continue
        if raw.startswith(" ") or raw.startswith("\t"):
            key, _, value = raw.strip().partition(":")
        else:
            key, _, value = raw.partition(":")
        key = key.strip()
        value = value.strip().strip("\"'")
        if key and value and key != "metadata":
            meta[key] = value
    body = text[match.end() :]
    if body.startswith("\n"):
        body = body[1:]
    return meta, body


def parse_inline(text: str) -> list[tuple[str, str]]:
    """拆成 (text|bold|code|bold_code, 内容)。先识别成对反引号，再识别成对 **。"""
    out: list[tuple[str, str]] = []
    buf: list[str] = []
    i = 0
    n = len(text)

    def flush() -> None:
        if buf:
            out.append(("text", "".join(buf)))
            buf.clear()

    while i < n:
        if text.startswith("`", i):
            j = text.find("`", i + 1)
            if j == -1:
                buf.append(text[i:])
                break
            flush()
            out.append(("code", text[i + 1 : j]))
            i = j + 1
            continue
        if text.startswith("**", i):
            j = text.find("**", i + 2)
            if j == -1:
                buf.append(text[i:])
                break
            flush()
            inner = text[i + 2 : j]
            if "`" in inner:
                k = 0
                while k < len(inner):
                    if inner.startswith("`", k):
                        m = inner.find("`", k + 1)
                        if m == -1:
                            out.append(("bold", inner[k:]))
                            break
                        out.append(("bold_code", inner[k + 1 : m]))
                        k = m + 1
                    else:
                        m = inner.find("`", k)
                        if m == -1:
                            out.append(("bold", inner[k:]))
                            break
                        if m > k:
                            out.append(("bold", inner[k:m]))
                        k = m
            elif inner:
                out.append(("bold", inner))
            i = j + 2
            continue
        buf.append(text[i])
        i += 1
    flush()
    return [(kind, value) for kind, value in out if value]


def set_run_font(
    run,
    *,
    east_asia: str,
    ascii_name: str,
    size_pt: float,
    bold: bool = False,
    color: str | None = None,
    italic: bool = False,
) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    run.font.name = ascii_name
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_name)
    rFonts.set(qn("w:hAnsi"), ascii_name)
    rFonts.set(qn("w:cs"), ascii_name)
    rFonts.set(qn("w:eastAsia"), east_asia)


def set_style_font(style, east_asia: str, ascii_name: str, size_pt: float, bold: bool | None, color: str | None) -> None:
    style.font.size = Pt(size_pt)
    style.font.name = ascii_name
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_name)
    rFonts.set(qn("w:hAnsi"), ascii_name)
    rFonts.set(qn("w:cs"), ascii_name)
    rFonts.set(qn("w:eastAsia"), east_asia)


def shade_paragraph(paragraph, fill: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def border_paragraph(paragraph, *, left: str | None = None, bottom: str | None = None) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    if left:
        el = OxmlElement("w:left")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "10")
        el.set(qn("w:color"), left)
        pBdr.append(el)
    if bottom:
        el = OxmlElement("w:bottom")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "1")
        el.set(qn("w:color"), bottom)
        pBdr.append(el)
    pPr.append(pBdr)


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, dxa: int = 60) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(dxa))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def set_table_borders(table, color: str = "C5D0DC") -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_table_width(table, width_dxa: int) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(width_dxa))
    tblW.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def mark_header_row(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    trPr.append(header)


def add_field(paragraph, instruction: str, *, east_asia: str, ascii_name: str, size_pt: float, color: str | None) -> None:
    run = paragraph.add_run()
    set_run_font(run, east_asia=east_asia, ascii_name=ascii_name, size_pt=size_pt, color=color)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    # 打开前的占位，避免页码域在未更新时完全空白
    placeholder = OxmlElement("w:t")
    placeholder.text = "1" if "PAGE" in instruction and "NUMPAGES" not in instruction else ""
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(begin)
    r.append(instr)
    r.append(separate)
    r.append(placeholder)
    r.append(end)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, "宋体", "Times New Roman", 12, False, "222222")
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    specs = [
        ("Heading 1", "黑体", 16, True, NAVY, 16, 8),
        ("Heading 2", "黑体", 14, True, NAVY, 12, 4),
        ("Heading 3", "黑体", 12, True, "24364A", 10, 2),
        ("Heading 4", "黑体", 12, True, "24364A", 8, 2),
    ]
    for name, east, size, bold, color, before, after in specs:
        style = doc.styles[name]
        set_style_font(style, east, "Times New Roman", size, bold, color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.page_break_before = True


def configure_section(section, content_width_dxa: int) -> int:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.6)
    section.different_first_page_header_footer = True
    return content_width_dxa


def is_cjk_label(text: str) -> bool:
    cjk = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")
    return cjk > 0 and cjk >= len(text.replace(" ", "")) / 2


def add_runs(
    paragraph,
    text: str,
    *,
    size: float = 12,
    east_asia: str = "宋体",
    ascii_name: str = "Times New Roman",
    color: str | None = None,
    base_bold: bool = False,
) -> None:
    for kind, value in parse_inline(text):
        run = paragraph.add_run(value)
        if kind in {"code", "bold_code"} and is_cjk_label(value):
            set_run_font(
                run,
                east_asia="黑体",
                ascii_name="Times New Roman",
                size_pt=size,
                bold=True,
                color=color or NAVY,
            )
        elif kind == "code":
            set_run_font(
                run,
                east_asia="宋体",
                ascii_name="Consolas",
                size_pt=max(size - 1, 9),
                bold=base_bold,
                color="9A3412",
            )
        elif kind == "bold_code":
            set_run_font(
                run,
                east_asia="宋体",
                ascii_name="Consolas",
                size_pt=max(size - 1, 9),
                bold=True,
                color="9A3412",
            )
        elif kind == "bold":
            set_run_font(
                run,
                east_asia=east_asia,
                ascii_name=ascii_name,
                size_pt=size,
                bold=True,
                color=color,
            )
        else:
            set_run_font(
                run,
                east_asia=east_asia,
                ascii_name=ascii_name,
                size_pt=size,
                bold=base_bold,
                color=color,
            )


def is_table_separator(line: str) -> bool:
    cells = split_row(line)
    return bool(cells) and all(SEP_CELL_RE.match(cell.replace(" ", "")) for cell in cells)


def split_row(line: str) -> list[str]:
    raw = line.strip()
    if not raw.startswith("|"):
        return []
    if raw.endswith("|"):
        raw = raw[1:-1]
    else:
        raw = raw[1:]
    return [cell.strip() for cell in raw.split("|")]


def looks_like_table_start(lines: list[str], index: int) -> bool:
    if not lines[index].lstrip().startswith("|"):
        return False
    if index + 1 >= len(lines):
        return False
    return is_table_separator(lines[index + 1])


def add_table(doc: Document, rows: list[list[str]], width_dxa: int) -> None:
    cols = max(len(row) for row in rows)
    normalized = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=cols)
    table.autofit = False
    table.allow_autofit = False
    set_table_width(table, width_dxa)
    set_table_borders(table)
    col_w = width_dxa // cols
    font_size = 8 if cols >= 7 else 9 if cols >= 5 else 10.5
    for r_idx, row in enumerate(normalized):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = Twips(col_w)
            set_cell_margins(cell, 50)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(col_w))
            tcW.set(qn("w:type"), "dxa")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
                add_runs(p, text, size=font_size, east_asia="黑体", base_bold=True, color="FFFFFF")
            else:
                if r_idx % 2 == 0:
                    set_cell_shading(cell, ZEBRA)
                add_runs(p, text, size=font_size)
        if r_idx == 0:
            mark_header_row(table.rows[r_idx])
        else:
            prevent_row_split(table.rows[r_idx])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(2)
    spacer.paragraph_format.space_after = Pt(2)


def add_meta_table(doc: Document, meta: dict[str, str], path: str, width_dxa: int) -> None:
    labels = [
        ("name", "name"),
        ("中文名称", "中文名称"),
        ("description", "用途"),
        ("所属分支", "所属分支"),
        ("状态", "状态"),
    ]
    rows = [(label, meta[key]) for key, label in labels if meta.get(key)]
    rows.append(("文件", path))
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    set_table_width(table, width_dxa)
    set_table_borders(table, "D5DEE8")
    label_w = 1800
    value_w = width_dxa - label_w
    for i, (label, value) in enumerate(rows):
        left, right = table.cell(i, 0), table.cell(i, 1)
        left.width = Twips(label_w)
        right.width = Twips(value_w)
        set_cell_shading(left, LABEL_BG)
        set_cell_margins(left, 60)
        set_cell_margins(right, 60)
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lp.paragraph_format.space_before = Pt(1)
        lp.paragraph_format.space_after = Pt(1)
        rp.paragraph_format.space_before = Pt(1)
        rp.paragraph_format.space_after = Pt(1)
        add_runs(lp, label, size=10.5, east_asia="黑体", base_bold=True, color=NAVY)
        add_runs(rp, value, size=10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(2)
    spacer.paragraph_format.space_after = Pt(4)


def render_markdown(doc: Document, body: str, width_dxa: int, *, skip_first_h1: str | None) -> None:
    lines = body.splitlines()
    i = 0
    skipped_h1 = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            block: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            if not block:
                block = [""]
            for idx, code_line in enumerate(block):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2 if idx == 0 else 0)
                p.paragraph_format.space_after = Pt(2 if idx == len(block) - 1 else 0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Cm(0.3)
                shade_paragraph(p, CODE_BG)
                run = p.add_run(code_line if code_line else " ")
                set_run_font(run, east_asia="宋体", ascii_name="Consolas", size_pt=9, color="1F2937")
            continue

        if looks_like_table_start(lines, i):
            rows = [split_row(lines[i])]
            i += 2  # header + separator
            while i < len(lines) and lines[i].lstrip().startswith("|") and lines[i].strip():
                if is_table_separator(lines[i]):
                    i += 1
                    continue
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows, width_dxa)
            continue

        if stripped.startswith(">"):
            while i < len(lines) and lines[i].strip().startswith(">"):
                content = re.sub(r"^\s*>\s?", "", lines[i])
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.45)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                shade_paragraph(p, QUOTE_BG)
                border_paragraph(p, left=NAVY)
                if content:
                    add_runs(p, content, size=10.5, color="1F2937")
                else:
                    run = p.add_run(" ")
                    set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size_pt=6)
                i += 1
            continue

        check = CHECK_RE.match(line)
        if check:
            indent, mark, content = check.groups()
            level = len(indent.replace("\t", "  ")) // 2
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(2)
            box = "☑" if mark.lower() == "x" else "☐"
            run = p.add_run(f"{box} ")
            set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size_pt=12)
            add_runs(p, content, size=12)
            i += 1
            continue

        listed = LIST_RE.match(line)
        if listed and not stripped.startswith("**"):
            indent, marker, content = listed.groups()
            level = len(indent.replace("\t", "  ")) // 2
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(2)
            bullet = marker if marker[0].isdigit() else "•"
            run = p.add_run(f"{bullet} ")
            set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size_pt=12, bold=False, color=NAVY)
            add_runs(p, content, size=12)
            i += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            hashes, title = heading.groups()
            level = len(hashes)
            if level == 1 and skip_first_h1 and not skipped_h1:
                skipped_h1 = True
                if title != skip_first_h1:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(6)
                    add_runs(p, title, size=12, east_asia="黑体", color=NAVY, base_bold=True)
                i += 1
                continue
            style = f"Heading {min(level, 4)}"
            p = doc.add_paragraph(style=style)
            add_runs(p, title, size={1: 16, 2: 14, 3: 12, 4: 12}[min(level, 4)], east_asia="黑体", color=NAVY, base_bold=True)
            i += 1
            continue

        if stripped in {"---", "***"}:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            border_paragraph(p, bottom=RULE)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped, size=12)
        i += 1


def add_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    border_paragraph(hp, bottom=NAVY)
    run = hp.add_run("创造性答复撰写 Skill 汇编")
    set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size_pt=9, color="64748B")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(2)
    add_runs(fp, "第 ", size=9, color="64748B")
    add_field(fp, " PAGE ", east_asia="宋体", ascii_name="Times New Roman", size_pt=9, color="64748B")
    add_runs(fp, " 页 / 共 ", size=9, color="64748B")
    add_field(fp, " NUMPAGES ", east_asia="宋体", ascii_name="Times New Roman", size_pt=9, color="64748B")
    add_runs(fp, " 页", size=9, color="64748B")

    # 封面不显示页眉
    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    ff = first_footer.paragraphs[0]
    ff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(ff, "审查意见答复 · 创造性答复", size=9, color="94A3B8")


def build_cover(doc: Document, width_dxa: int) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("创造性答复撰写")
    set_run_font(run, east_asia="黑体", ascii_name="Times New Roman", size_pt=26, bold=True, color=NAVY)

    title2 = doc.add_paragraph()
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2.paragraph_format.space_before = Pt(0)
    title2.paragraph_format.space_after = Pt(8)
    run = title2.add_run("Skill 全文汇编")
    set_run_font(run, east_asia="黑体", ascii_name="Times New Roman", size_pt=26, bold=True, color=NAVY)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_after = Pt(10)
    border_paragraph(line, bottom=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(14)
    run = sub.add_run("审查意见答复　·　专利法第22条第3款")
    set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size_pt=12, color="475569")

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    intro.paragraph_format.first_line_indent = Cm(0.74)
    add_runs(
        intro,
        "本文收录仓库「审查意见答复」分支现有全部 skill 正文，按跑案顺序编排。"
        "当前只有创造性答复这一套；新颖性、修改超范围、不清楚或得不到说明书支持的答复尚未单列。"
        "跑案入口是「创造性答复总控」。各节正文与对应 SKILL.md 一致，并保留 name、用途与状态。",
    )

    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(6)
    run = h.add_run("目录")
    set_run_font(run, east_asia="黑体", ascii_name="Times New Roman", size_pt=14, bold=True, color=NAVY)

    rows = [["序号", "名称", "角色"]]
    for idx, section in enumerate(SECTIONS, start=1):
        rows.append([cn_index(idx), section["title"], section["role"]])
    add_table(doc, rows, width_dxa)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    add_runs(
        note,
        "说明两节在前，其后 9 个 skill 按总控挂接顺序排列。旧名「创造性答复意见陈述书撰写」仅作跳转，不单独成文。",
        size=10.5,
        color="475569",
    )


def assert_manifest_complete() -> None:
    found = {p.relative_to(REPLY_ROOT).as_posix() for p in REPLY_ROOT.rglob("SKILL.md")}
    listed = {item["path"] for item in SECTIONS if item["path"].endswith("SKILL.md")}
    missing = found - listed
    extra = listed - found
    if missing or extra:
        raise SystemExit(f"skill 清单与目录不一致。缺少：{sorted(missing)}；多余：{sorted(extra)}")
    for item in SECTIONS:
        path = REPLY_ROOT / item["path"]
        if not path.is_file():
            raise SystemExit(f"找不到文件：{path}")


def build_document() -> Document:
    assert_manifest_complete()
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    # A4 内容区宽度（dxa）：11906 - 左右边距。1.9cm ≈ 1077 dxa
    margin_dxa = int(1.9 / 2.54 * 1440)
    content_width = 11906 - margin_dxa * 2
    configure_section(section, content_width)
    add_header_footer(section)

    core = doc.core_properties
    core.title = "创造性答复撰写 Skill 全文汇编"
    core.subject = "审查意见答复 · 创造性答复"
    core.category = "专利 skill"
    core.comments = "由 scripts/export_reply_skills_docx.py 自 SKILL.md 生成"

    build_cover(doc, content_width)

    for idx, item in enumerate(SECTIONS, start=1):
        path = REPLY_ROOT / item["path"]
        text = path.read_text(encoding="utf-8")
        meta, body = parse_frontmatter(text)
        heading = doc.add_paragraph(style="Heading 1")
        add_runs(
            heading,
            f"{cn_index(idx)}、{item['title']}",
            size=16,
            east_asia="黑体",
            color=NAVY,
            base_bold=True,
        )
        role = doc.add_paragraph()
        role.paragraph_format.space_before = Pt(0)
        role.paragraph_format.space_after = Pt(6)
        add_runs(role, item["role"], size=10.5, color="64748B")
        if meta:
            add_meta_table(doc, meta, ".cursor/skills/专利/审查意见答复/" + item["path"], content_width)
        else:
            src = doc.add_paragraph()
            src.paragraph_format.space_after = Pt(6)
            add_runs(src, "来源：" + ".cursor/skills/专利/审查意见答复/" + item["path"], size=10.5, color="64748B")
        render_markdown(doc, body, content_width, skip_first_h1=item["title"])
    return doc


def plain_text(doc: Document) -> str:
    chunks: list[str] = []
    for p in doc.paragraphs:
        if p.text:
            chunks.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    chunks.append(cell.text)
    return "\n".join(chunks)


def rendered_plain(text: str) -> str:
    return "".join(value for _, value in parse_inline(text))


def significant_fragments(source: str) -> list[str]:
    """按导出后的可见文字抽查，星号与反引号的处理与正文渲染一致。"""
    meta, body = parse_frontmatter(source)
    fragments: list[str] = []
    for value in meta.values():
        if len(value) >= 8:
            fragments.append(value[:40])
    in_code = False
    for raw in body.splitlines():
        line = raw.strip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if not line or is_table_separator(line):
            continue
        if in_code:
            if len(line) >= 8:
                fragments.append(line[:48])
            continue
        if line.startswith("|"):
            for cell in split_row(line):
                visible = rendered_plain(cell).strip()
                if len(visible) >= 6:
                    fragments.append(visible[:36])
            continue
        line = re.sub(r"^#{1,4}\s+", "", line)
        line = re.sub(r"^>\s?", "", line)
        line = re.sub(r"^[-*+]\s+\[[ xX]\]\s+", "", line)
        line = re.sub(r"^[-*+]\s+", "", line)
        line = re.sub(r"^\d+\.\s+", "", line)
        visible = rendered_plain(line).strip()
        if len(visible) >= 8:
            fragments.append(visible[:48])
    return fragments


def verify(doc: Document) -> None:
    text = plain_text(doc)
    missing: list[str] = []
    checked = 0
    for item in SECTIONS:
        source = (REPLY_ROOT / item["path"]).read_text(encoding="utf-8")
        for fragment in significant_fragments(source):
            checked += 1
            if fragment not in text:
                missing.append(f"{item['title']}: {fragment}")
    skill_count = sum(1 for item in SECTIONS if item["kind"] == "skill")
    if skill_count != 9:
        raise SystemExit(f"期望 9 个 skill，实际 {skill_count}")
    if missing:
        preview = "\n".join(missing[:30])
        raise SystemExit(f"导出后缺 {len(missing)} 处原文片段（抽查 {checked} 处）：\n{preview}")
    print(f"核对通过：{len(SECTIONS)} 节，其中 skill {skill_count} 个；抽查片段 {checked} 处均在文档中。")
    print(f"段落 {len(doc.paragraphs)}，表格 {len(doc.tables)}。")


def self_check_inline() -> None:
    samples = {
        "只交一份并用 `**…**` 标加粗": "只交一份并用 **…** 标加粗",
        "**`创造性答复总控`**": "创造性答复总控",
        "关键句已`**加粗**`": "关键句已**加粗**",
        "须「逐条 **3～5 条**」": "须「逐条 3～5 条」",
    }
    for src, expected in samples.items():
        got = rendered_plain(src)
        if got != expected:
            raise SystemExit(f"行内标记解析不符：{src!r} -> {got!r}，期望 {expected!r}")


def main() -> int:
    parser = argparse.ArgumentParser(description="导出审查意见答复全部 skill 为 Word")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    self_check_inline()
    doc = build_document()
    verify(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"已写入 {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
